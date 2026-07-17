#!/usr/bin/env python3
"""Build an auditable workbook for the local IWG/RG-CMA experiments."""

from __future__ import annotations

import argparse
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[2]
EXPERIMENTS = ROOT / "outputs/agentguard/experiments"
DEFAULT_OUTPUT = ROOT / "outputs/agentguard/reports/IWG_RG_CMA实验汇总.xlsx"
METRICS = ("HOTA", "MOTA", "IDF1", "DetA", "AssA")


@dataclass(frozen=True)
class Experiment:
    dataset: str
    name: str
    directory: str
    train_data: str
    event_source: str
    runtime_kf: str
    schedule: str
    memory_shards: int | None
    epochs_per_shard: int | None
    shard_cycles: int | None
    effective_full_epochs: float | None
    batch_size: int | None
    lr: float | None
    weight_decay: float | None
    seed: int | None

    @property
    def root(self) -> Path:
        return EXPERIMENTS / self.directory if self.directory else ROOT


HEADERS = [
    "数据集",
    "实验简称",
    "实验目录",
    "训练数据",
    "事件/标签来源",
    "运行时运动模型",
    "训练方式",
    "memory_shards",
    "epochs_per_shard",
    "shard_cycles",
    "等效全量epoch",
    "batch_size",
    "lr",
    "weight_decay",
    "seed",
    "权重文件",
    "Epoch",
    "Gate",
    "处理",
    "评测集",
    "HOTA",
    "MOTA",
    "IDF1",
    "DetA",
    "AssA",
    "指标来源",
    "备注",
]


def rel(path: Path | str | None) -> str:
    if not path:
        return ""
    path = Path(path)
    try:
        return str(path.resolve().relative_to(ROOT))
    except ValueError:
        return str(path)


def load_json(path: Path) -> dict[str, Any]:
    if not path.is_file():
        raise FileNotFoundError(path)
    with path.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def parse_metric_log(path: Path) -> dict[str, float]:
    text = path.read_text(encoding="utf-8", errors="replace")
    match = re.search(
        r"^HOTA\s+MOTA\s+IDF1\s+DetA\s+AssA\s*$\s*"
        r"([0-9.]+)\s+([0-9.]+)\s+([0-9.]+)\s+([0-9.]+)\s+([0-9.]+)",
        text,
        flags=re.MULTILINE,
    )
    if not match:
        raise ValueError(f"No combined metric block found in {path}")
    return dict(zip(METRICS, map(float, match.groups())))


def checkpoint(exp: Experiment, epoch: int | str | None, legacy: str = "") -> str:
    if legacy:
        return rel(exp.root / legacy)
    if epoch is None or not exp.directory:
        return ""
    return rel(exp.root / "checkpoints" / f"iwg_rg_cma_epoch{int(epoch):03d}.pt")


def build_row(
    exp: Experiment,
    *,
    epoch: int | str | None,
    gate: str,
    processing: str,
    split: str,
    metrics: dict[str, float | None] | None,
    source: str,
    note: str = "",
    checkpoint_path: str = "",
) -> dict[str, Any]:
    values = metrics or {}
    return {
        "数据集": exp.dataset,
        "实验简称": exp.name,
        "实验目录": rel(exp.root) if exp.directory else "",
        "训练数据": exp.train_data,
        "事件/标签来源": exp.event_source,
        "运行时运动模型": exp.runtime_kf,
        "训练方式": exp.schedule,
        "memory_shards": exp.memory_shards,
        "epochs_per_shard": exp.epochs_per_shard,
        "shard_cycles": exp.shard_cycles,
        "等效全量epoch": exp.effective_full_epochs,
        "batch_size": exp.batch_size,
        "lr": exp.lr,
        "weight_decay": exp.weight_decay,
        "seed": exp.seed,
        "权重文件": checkpoint_path or checkpoint(exp, epoch),
        "Epoch": epoch,
        "Gate": gate,
        "处理": processing,
        "评测集": split,
        **{metric: values.get(metric) for metric in METRICS},
        "指标来源": source,
        "备注": note,
    }


def add_summary_json(
    rows: list[dict[str, Any]], exp: Experiment, filename: str = "evaluation_summary.json"
) -> None:
    path = exp.root / filename
    data = load_json(path)
    for epoch_text, gates in sorted(data["cases"].items(), key=lambda item: int(item[0])):
        epoch = int(epoch_text)
        for gate in ("base", "final"):
            for processing in ("raw", "post"):
                rows.append(
                    build_row(
                        exp,
                        epoch=epoch,
                        gate=gate,
                        processing=processing,
                        split="all" if exp.dataset == "MOT20" else "val",
                        metrics=gates[gate][processing],
                        source=rel(path),
                    )
                )


def add_mot17_log_set(
    rows: list[dict[str, Any]], exp: Experiment, *, epoch: int, prefix: str
) -> None:
    for gate in ("base", "final"):
        for processing in ("raw", "post"):
            path = exp.root / "logs" / f"{prefix}_{gate}_{processing}.log"
            rows.append(
                build_row(
                    exp,
                    epoch=epoch,
                    gate=gate,
                    processing=processing,
                    split="all",
                    metrics=parse_metric_log(path),
                    source=rel(path),
                )
            )


def collect_rows() -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []

    baseline_mot17 = Experiment(
        "MOT17", "NSA Kalman Filter baseline", "", "无训练", "无", "NSA KF",
        "baseline", None, None, None, None, None, None, None, None,
    )
    baseline_mot20 = Experiment(
        "MOT20", "NSA Kalman Filter baseline", "", "无训练", "无", "NSA KF",
        "baseline", None, None, None, None, None, None, None, None,
    )
    baseline_sports = Experiment(
        "SportsMOT", "NSA Kalman Filter baseline", "", "无训练", "无", "NSA KF",
        "baseline", None, None, None, None, None, None, None, None,
    )

    for processing, values in {
        "raw": dict(zip(METRICS, (0.770290, 0.885830, 0.874057, 0.773059, 0.771415))),
        "post": dict(zip(METRICS, (0.774487, 0.895625, 0.877821, 0.778875, 0.774544))),
    }.items():
        rows.append(build_row(
            baseline_mot17, epoch=None, gate="off", processing=processing, split="all",
            metrics=values, source="2026-07-16 对已有 mot17_all tracker 结果重跑 TrackEval",
            note="仅七个 MOT17-FRCNN train 序列。",
        ))

    old_iwg = Experiment(
        "MOT17", "旧 IWG（无 RG-CMA）", "iwg_v2_a_only_100e_seed42",
        "MOT17 七个 FRCNN train 序列", "NSA 事件 + safe gate 标签", "NSA KF",
        "全量常驻", 1, 100, 1, 100, 1024, 1e-4, 1e-4, 42,
    )
    for processing in ("raw", "post"):
        path = old_iwg.root / "eval_mot17_all/logs" / f"{processing}.log"
        rows.append(build_row(
            old_iwg, epoch=100, gate="IWG", processing=processing, split="all",
            metrics=parse_metric_log(path), source=rel(path),
            checkpoint_path=checkpoint(old_iwg, None, "checkpoints/iwg/iwg_last.pt"),
        ))
    rows.append(build_row(
        old_iwg, epoch=100, gate="IWG", processing="raw", split="test",
        metrics={"HOTA": 0.6609},
        source="本地 ZIP 文件名：mot17_test_0.80_a_only_bs1024_100e_v2_agentguard_iwg_66.09.zip",
        note="官方 test 其余四项未保存在本地。",
        checkpoint_path=checkpoint(old_iwg, None, "checkpoints/iwg/iwg_last.pt"),
    ))
    rows.append(build_row(
        old_iwg, epoch=100, gate="IWG", processing="post", split="test",
        metrics=dict(zip(METRICS, (0.673490, 0.818488, 0.837289, 0.663682, 0.686532))),
        source="用户在会话中提供的官方 test 结果；本地 ZIP HOTA=67.34",
        checkpoint_path=checkpoint(old_iwg, None, "checkpoints/iwg/iwg_last.pt"),
    ))

    mot17_2048 = Experiment(
        "MOT17", "RG-CMA bs2048", "iwg_rg_cma_v1_trainall_seed42",
        "MOT17 七个 FRCNN train 序列", "NSA 事件 + safe-direct 标签", "NSA KF",
        "全量常驻", 1, 100, 1, 100, 2048, 1e-4, 1e-4, 42,
    )
    add_mot17_log_set(rows, mot17_2048, epoch=100, prefix="eval_iwg_rg_cma_v1_trainall_seed42")

    mot17_1024 = Experiment(
        "MOT17", "RG-CMA bs1024", "iwg_rg_cma_v1_trainall_seed42_bs1024_native_log",
        "MOT17 七个 FRCNN train 序列", "NSA 事件 + safe-direct 标签", "NSA KF",
        "全量常驻", 1, 100, 1, 100, 1024, 1e-4, 1e-4, 42,
    )
    add_mot17_log_set(
        rows, mot17_1024, epoch=100,
        prefix="eval_iwg_rg_cma_v1_trainall_seed42_bs1024_native_log",
    )
    rows.append(build_row(
        mot17_1024, epoch=100, gate="base", processing="post", split="test",
        metrics=dict(zip(METRICS, (0.668066, 0.811560, 0.824274, 0.659357, 0.680044))),
        source="用户在会话中提供的官方 test 结果",
    ))
    rows.append(build_row(
        mot17_1024, epoch=100, gate="final", processing="post", split="test",
        metrics=dict(zip(METRICS, (0.667370, 0.811651, 0.822698, 0.659406, 0.678575))),
        source="用户在会话中提供的官方 test 结果；本地 ZIP HOTA=66.73",
    ))

    mot20_transfer = Experiment(
        "MOT17", "MOT20旧方式e50直接迁移", "mot20_e050_transfer_mot17",
        "MOT20 train", "NSA 事件 + safe-direct 标签", "NSA KF",
        "MOT20旧方式训练；MOT17零微调", 5, 10, 2, 10, 1024, 1e-4, 1e-4, 42,
    )
    source_checkpoint = (
        EXPERIMENTS / "iwg_rg_cma_v1_mot20_trainall_seed42_bs1024_shard20x2/"
        "checkpoints/iwg_rg_cma_epoch050.pt"
    )
    for gate in ("base", "final"):
        for processing in ("raw", "post"):
            path = mot20_transfer.root / f"eval_{gate}_{processing}.log"
            rows.append(build_row(
                mot20_transfer, epoch=50, gate=gate, processing=processing, split="all",
                metrics=parse_metric_log(path), source=rel(path),
                checkpoint_path=rel(source_checkpoint),
                note="在 MOT17/all 上直接使用 MOT20 权重。",
            ))
    rows.append(build_row(
        mot20_transfer, epoch=50, gate="final", processing="post", split="test",
        metrics={"HOTA": 0.6720},
        source="本地 ZIP 文件名：MOT17_test_mot20_old_e050_xfer_final_post_67.2.zip",
        note="官方 test 其余四项未保存在本地。",
        checkpoint_path=rel(source_checkpoint),
    ))

    mot17_finetune = Experiment(
        "MOT17", "MOT20 e50 -> MOT17续训25e（最佳权重）",
        "iwg_rg_cma_mot20e050_to_mot17_finetune25_cachefix_seed42_bs1024",
        "MOT17 七个 FRCNN train 序列", "NSA 事件 + safe-direct 标签", "NSA KF",
        "MOT20旧方式e50 warm-start，MOT17全量续训", 1, 25, 1, 25,
        1024, 1e-5, 1e-4, 42,
    )
    rows.append(build_row(
        mot17_finetune, epoch=25, gate="final", processing="post", split="test",
        metrics=None,
        source="本地存在提交 ZIP，但文件名与日志均未保存官方分数",
        note="用户认定的 MOT17 最佳权重；五项官方 test 指标待从提交平台补录。",
    ))

    mot20_old = Experiment(
        "MOT20", "RG-CMA旧方式", "iwg_rg_cma_v1_mot20_trainall_seed42_bs1024_shard20x2",
        "MOT20 train 全四序列", "NSA 事件 + safe-direct 标签", "NSA KF",
        "旧：同一20%分片连续10轮", 5, 10, 2, 20, 1024, 1e-4, 1e-4, 42,
    )
    old_eval_path = mot20_old.root / "evaluation.json"
    old_eval = load_json(old_eval_path)["cases"]
    for key in ("baseline_raw", "baseline_post"):
        processing = key.rsplit("_", 1)[1]
        rows.append(build_row(
            baseline_mot20, epoch=None, gate="off", processing=processing, split="all",
            metrics=old_eval[key]["combined"], source=rel(old_eval_path),
        ))
    for gate in ("base", "final"):
        for processing in ("raw", "post"):
            rows.append(build_row(
                mot20_old, epoch=100, gate=gate, processing=processing, split="all",
                metrics=old_eval[f"{gate}_{processing}"]["combined"], source=rel(old_eval_path),
                note="evaluation.json 对应该实验最终（epoch100）checkpoint。",
            ))
    rows.append(build_row(
        mot20_old, epoch=50, gate="final", processing="post", split="test",
        metrics={"HOTA": 0.6625},
        source="本地 ZIP 文件名：mot20_iwg_rg_cma_epoch050_final_post_66.25.zip",
        note="MOT20 当前最佳 test 权重；其余四项未保存在本地。",
    ))

    mot20_new = Experiment(
        "MOT20", "RG-CMA新方式", "iwg_rg_cma_v1_mot20_interleaved_shard1x20_seed42_bs1024",
        "MOT20 train 全四序列", "NSA 事件 + safe-direct 标签", "NSA KF",
        "新：每轮切换20%分片", 5, 1, 20, 20, 1024, 1e-4, 1e-4, 42,
    )
    add_summary_json(rows, mot20_new)
    rows.append(build_row(
        mot20_new, epoch=100, gate="final", processing="post", split="test",
        metrics={"HOTA": 0.6612},
        source="本地 ZIP 文件名：mot20_test_0.80_best_interleaved_e100_iwg_attn_final_post_66.12.zip",
        note="官方 test 其余四项未保存在本地。",
    ))

    mamba_native = Experiment(
        "MOT20", "Mamba-native事件训练RG-CMA", "iwg_rg_cma_mamba_native_mot20_seed42_bs1024_shard5x2",
        "MOT20 train 全四序列", "Mamba KF 导出事件 + mamba_native motion target", "NSA KF",
        "旧：同一20%分片连续10轮", 5, 10, 2, 20, 1024, 1e-4, 1e-4, 42,
    )
    mamba_results = {
        (50, "raw"): (0.787636, 0.934492, 0.914958, 0.806062, 0.770204),
        (50, "post"): (0.791342, 0.938477, 0.916034, 0.810256, 0.773464),
        (100, "raw"): (0.789268, 0.934761, 0.917063, 0.806180, 0.773275),
        (100, "post"): (0.792715, 0.938301, 0.918014, 0.810002, 0.776386),
    }
    for (epoch, processing), values in mamba_results.items():
        source = (
            rel(mamba_native.root / f"logs/eval_e{epoch:03d}_final_all_post.log")
            if processing == "post"
            else "2026-07-16 对已有 tracker 结果重跑 TrackEval"
        )
        rows.append(build_row(
            mamba_native, epoch=epoch, gate="final", processing=processing, split="all",
            metrics=dict(zip(METRICS, values)), source=source,
            note="训练事件来自 Mamba；在线 tracker 仍使用 NSA KF。",
        ))

    baseline_sports_raw = dict(zip(
        METRICS, (0.814892, 0.982590, 0.832692, 0.913213, 0.727347)
    ))
    rows.append(build_row(
        baseline_sports, epoch=None, gate="off", processing="raw", split="val",
        metrics=baseline_sports_raw,
        source="2026-07-16 对已有 sportsmot_val tracker 结果重跑 TrackEval",
        note="baseline post tracker 文件不完整，未写入 post 指标。",
    ))

    sports_old = Experiment(
        "SportsMOT", "RG-CMA旧方式 train-only", "iwg_rg_cma_v1_sportsmot_old_shard10x2_seed42_bs1024",
        "SportsMOT train", "NSA 事件 + safe-direct 标签", "NSA KF",
        "旧：同一20%分片连续10轮", 5, 10, 2, 20, 1024, 1e-4, 1e-4, 42,
    )
    add_summary_json(rows, sports_old)
    rows.append(build_row(
        sports_old, epoch=100, gate="final", processing="post", split="test",
        metrics={"HOTA": 0.7582},
        source="用户提供的官方 test HOTA；本地提交 ZIP 存在",
        note="官方 test 其余四项未保存在本地。",
    ))

    sports_new = Experiment(
        "SportsMOT", "RG-CMA新方式 train-only", "iwg_rg_cma_v1_sportsmot_interleaved_shard1x20_seed42_bs1024",
        "SportsMOT train", "NSA 事件 + safe-direct 标签", "NSA KF",
        "新：每轮切换20%分片", 5, 1, 20, 20, 1024, 1e-4, 1e-4, 42,
    )
    add_summary_json(rows, sports_new)

    sports_trainval = Experiment(
        "SportsMOT", "RG-CMA train+val 200e（最佳权重）",
        "iwg_rg_cma_v1_sportsmot_trainval_seed42_bs1024_shard10x2_200e",
        "SportsMOT train+val", "NSA 事件 + safe-direct 标签", "NSA KF",
        "旧：同一10%分片连续10轮", 10, 10, 2, 20, 1024, 1e-4, 1e-4, 42,
    )
    trainval_results = {
        (100, "raw"): (0.831535, 0.982509, 0.831499, 0.932985, 0.741310),
        (100, "post"): (0.832354, 0.983865, 0.832248, 0.933645, 0.742245),
        (200, "raw"): (0.833733, 0.982505, 0.834608, 0.932987, 0.745228),
        (200, "post"): (0.834579, 0.983825, 0.835382, 0.933673, 0.746191),
    }
    for (epoch, processing), values in trainval_results.items():
        rows.append(build_row(
            sports_trainval, epoch=epoch, gate="final", processing=processing, split="val",
            metrics=dict(zip(METRICS, values)),
            source="2026-07-16 对已有 tracker 结果重跑 TrackEval",
            note="训练已包含 val，因此该 val 指标仅用于回顾拟合，不是独立泛化验证。",
        ))
    rows.append(build_row(
        sports_trainval, epoch=200, gate="final", processing="post", split="test",
        metrics={"HOTA": 0.7589},
        source="本地 ZIP 文件名：sportsmot_test_0.80_trainval_200e_iwg_attn_final_post_75.89.zip",
        note="SportsMOT 当前最佳 test 权重；其余四项未保存在本地。",
    ))

    sports_trainval_new = Experiment(
        "SportsMOT", "RG-CMA train+val 新方式 200e",
        "iwg_rg_cma_v1_sportsmot_trainval_seed42_bs1024_interleaved_shard1x20_200e",
        "SportsMOT train+val", "NSA 事件 + safe-direct 标签", "NSA KF",
        "新：每轮切换10%分片", 10, 1, 20, 20, 1024, 1e-4, 1e-4, 42,
    )
    trainval_new_raw_results = {
        50: (0.831938, 0.982559, 0.831541, 0.933542, 0.741582),
        100: (0.832265, 0.982529, 0.832300, 0.933244, 0.742404),
        150: (0.832821, 0.982515, 0.832830, 0.933132, 0.743490),
        200: (0.831036, 0.982525, 0.830775, 0.933052, 0.740371),
    }
    for epoch, values in trainval_new_raw_results.items():
        rows.append(build_row(
            sports_trainval_new,
            epoch=epoch,
            gate="final",
            processing="raw",
            split="val",
            metrics=dict(zip(METRICS, values)),
            source=rel(sports_trainval_new.root / f"logs/eval_e{epoch:03d}_final_raw.log"),
            note=(
                "新方式 raw val 最佳候选；val 已参与训练，仅用于实验内 checkpoint 对比。"
                if epoch == 150
                else "val 已参与训练，仅用于实验内 checkpoint 对比。"
            ),
        ))

    return rows


def style_metric_sheet(ws, rows: list[dict[str, Any]], table_name: str) -> None:
    ws.append(HEADERS)
    for row in rows:
        ws.append([row.get(header) for header in HEADERS])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False

    post_fill = PatternFill("solid", fgColor="E2F0D9")
    raw_fill = PatternFill("solid", fgColor="DDEBF7")
    test_fill = PatternFill("solid", fgColor="FFF2CC")
    missing_fill = PatternFill("solid", fgColor="FCE4D6")
    metric_columns = {name: HEADERS.index(name) + 1 for name in METRICS}
    process_col = HEADERS.index("处理") + 1
    split_col = HEADERS.index("评测集") + 1
    note_col = HEADERS.index("备注") + 1
    source_col = HEADERS.index("指标来源") + 1

    for row_index in range(2, ws.max_row + 1):
        processing = ws.cell(row_index, process_col).value
        split = ws.cell(row_index, split_col).value
        row_fill = post_fill if processing == "post" else raw_fill
        for col in range(1, ws.max_column + 1):
            ws.cell(row_index, col).fill = test_fill if split == "test" else row_fill
            ws.cell(row_index, col).alignment = Alignment(vertical="top", wrap_text=col in (3, 4, 5, 6, 7, 16, source_col, note_col))
        for col in metric_columns.values():
            cell = ws.cell(row_index, col)
            cell.number_format = "0.000000"
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if cell.value is None:
                cell.fill = missing_fill
        for name in ("lr", "weight_decay"):
            ws.cell(row_index, HEADERS.index(name) + 1).number_format = "0.00E+00"

    if ws.max_row >= 2:
        table = Table(displayName=table_name, ref=ws.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
            showRowStripes=False, showColumnStripes=False,
        )
        ws.add_table(table)
        hota_letter = get_column_letter(metric_columns["HOTA"])
        ws.conditional_formatting.add(
            f"{hota_letter}2:{hota_letter}{ws.max_row}",
            ColorScaleRule(
                start_type="min", start_color="F8696B",
                mid_type="percentile", mid_value=50, mid_color="FFEB84",
                end_type="max", end_color="63BE7B",
            ),
        )

    widths = {
        1: 13, 2: 32, 3: 55, 4: 28, 5: 34, 6: 18, 7: 28,
        8: 14, 9: 18, 10: 14, 11: 16, 12: 12, 13: 12, 14: 14,
        15: 10, 16: 68, 17: 10, 18: 11, 19: 10, 20: 10,
        21: 13, 22: 13, 23: 13, 24: 13, 25: 13, 26: 55, 27: 52,
    }
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def best_rows(all_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    selected: list[dict[str, Any]] = []
    predicates = (
        lambda row: row["实验简称"] == "MOT20 e50 -> MOT17续训25e（最佳权重）",
        lambda row: row["实验简称"] == "RG-CMA旧方式" and row["Epoch"] == 50 and row["评测集"] == "test",
        lambda row: row["实验简称"] == "RG-CMA train+val 200e（最佳权重）" and row["Epoch"] == 200,
    )
    for predicate in predicates:
        selected.extend(row for row in all_rows if predicate(row))
    return selected


def build_notes_sheet(ws, rows: list[dict[str, Any]]) -> None:
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 115
    notes = [
        ("工作簿用途", "汇总本机能追溯到的 IWG / RG-CMA combined 指标；不是论文最终表。"),
        ("数值单位", "HOTA/MOTA/IDF1/DetA/AssA 全部使用 0-1，小数保留六位。"),
        ("all", "MOT17 all 仅七个 FRCNN train 序列；MOT20 all 为四个 train 序列。"),
        ("val", "SportsMOT train-only 实验在 val 上评测；train+val 实验的 val 已参与训练，只能观察拟合。2026-07-16 起 validation 只记录 raw，不再运行 post。"),
        ("test", "test 指标来自官方评测或用户提供。若本地只保存 HOTA，其他四项留空，不以 0 代替。"),
        ("base gate", "Safe-Direct IWG 输出，不使用 RG-CMA correction。"),
        ("final gate", "base gate 加 RG-CMA 有界 correction 后的输出。"),
        ("raw", "tracker 原始输出，未做数据集对应的后处理。"),
        ("post", "MOT 使用 GSI，SportsMOT 使用线性插值；与 run.py --use_post 一致。"),
        ("旧分片方式", "epochs_per_shard=10：一个内存分片连续训练10轮后再切换。"),
        ("新分片方式", "epochs_per_shard=1：每轮切换分片，使学习率阶段在各时间段上更均衡。"),
        ("shard20x2 / shard10x2", "目录名是历史命名，判断方式以 memory_shards / epochs_per_shard / shard_cycles 三列的真实值为准。"),
        ("Mamba-native", "只用 Mamba KF 导出的事件和运动目标训练；评测时 tracker 仍使用 NSA KF。"),
        ("MOT17最佳权重", "MOT20旧方式 epoch50 warm-start 后，在 MOT17 全量续训25轮；本地提交 ZIP 没有官方分数。"),
        ("MOT20最佳权重", "旧方式 epoch50；已知 test final+post HOTA=0.6625。"),
        ("SportsMOT最佳权重", "train+val 旧分片方式 epoch200；已知 test final+post HOTA=0.7589。"),
        ("生成脚本", rel(Path(__file__))),
        ("生成行数", str(len(rows))),
    ]
    for index, (key, value) in enumerate(notes, start=1):
        ws.cell(index, 1, key)
        ws.cell(index, 2, value)
        ws.cell(index, 1).font = Font(bold=True, color="FFFFFF")
        ws.cell(index, 1).fill = PatternFill("solid", fgColor="1F4E78")
        ws.cell(index, 1).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(index, 2).alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[index].height = 31
    ws.freeze_panes = "A1"


def build_workbook(output: Path) -> None:
    rows = collect_rows()
    workbook = Workbook()
    workbook.remove(workbook.active)

    best = workbook.create_sheet("最佳权重")
    style_metric_sheet(best, best_rows(rows), "BestWeights")
    for dataset, table_name in (
        ("MOT17", "MOT17Metrics"),
        ("MOT20", "MOT20Metrics"),
        ("SportsMOT", "SportsMOTMetrics"),
    ):
        sheet = workbook.create_sheet(dataset)
        style_metric_sheet(sheet, [row for row in rows if row["数据集"] == dataset], table_name)
    all_metrics = workbook.create_sheet("全部指标")
    style_metric_sheet(all_metrics, rows, "AllMetrics")
    notes = workbook.create_sheet("配置说明")
    build_notes_sheet(notes, rows)

    workbook.calculation.fullCalcOnLoad = True
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    print(json.dumps({
        "output": str(output),
        "rows": len(rows),
        "by_dataset": {
            dataset: sum(row["数据集"] == dataset for row in rows)
            for dataset in ("MOT17", "MOT20", "SportsMOT")
        },
        "sheets": workbook.sheetnames,
    }, ensure_ascii=False, indent=2))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    build_workbook(args.output.resolve())


if __name__ == "__main__":
    main()
